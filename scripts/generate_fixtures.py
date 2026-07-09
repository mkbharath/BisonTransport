"""Generate test fixture files (PDF, Excel, Word, PNG) for all scenarios."""
import os

output_dir = os.path.join(os.path.dirname(__file__), "..", "test-emails", "fixtures")
os.makedirs(output_dir, exist_ok=True)


def create_pdf(filename, content):
    """Create a simple PDF with text content."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import inch

    path = os.path.join(output_dir, filename)
    c = canvas.Canvas(path, pagesize=letter)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(1 * inch, 10 * inch, "SHIPMENT REQUEST")
    c.setFont("Helvetica", 10)
    y = 9.5 * inch
    for line in content.split("\n"):
        c.drawString(1 * inch, y, line)
        y -= 14
        if y < 1 * inch:
            c.showPage()
            c.setFont("Helvetica", 10)
            y = 10 * inch
    c.save()
    print(f"  Created: {filename}")


def create_excel(filename, data):
    """Create an Excel file with order data."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    path = os.path.join(output_dir, filename)
    wb = Workbook()
    ws = wb.active
    ws.title = "Shipment Request"

    # Header
    headers = list(data[0].keys())
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        cell.font = Font(bold=True, color="FFFFFF")

    # Data rows
    for row_idx, row_data in enumerate(data, 2):
        for col, h in enumerate(headers, 1):
            ws.cell(row=row_idx, column=col, value=row_data[h])

    # Auto-width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 30)

    wb.save(path)
    print(f"  Created: {filename}")


def create_word(filename, content):
    """Create a Word document with order content."""
    from docx import Document
    from docx.shared import Pt

    path = os.path.join(output_dir, filename)
    doc = Document()
    doc.add_heading("SHIPMENT REQUEST", level=1)
    for line in content.split("\n"):
        if line.strip():
            doc.add_paragraph(line)
    doc.save(path)
    print(f"  Created: {filename}")


def create_image(filename, content):
    """Create a PNG image with order text (simulates scanned document)."""
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (800, 1000), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 14)
        font_bold = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 18)
    except:
        font = ImageFont.load_default()
        font_bold = font

    draw.text((50, 30), "SHIPMENT REQUEST", fill="black", font=font_bold)
    y = 70
    for line in content.split("\n"):
        draw.text((50, y), line, fill="black", font=font)
        y += 20

    path = os.path.join(output_dir, filename)
    img.save(path)
    print(f"  Created: {filename}")


# ═══════════════════════════════════════════════════════════════
# SCENARIOS
# ═══════════════════════════════════════════════════════════════

print("Generating test fixtures...")
print("=" * 50)

# --- Scenario 1: Perfect Order (PDF) ---
print("\n1. Perfect Order (PDF)")
create_pdf("scenario1-techvista-perfect.pdf", """
Customer: TechVista Solutions Inc.
Contact: Michael Chen
Email: bharathm@ideyalabs.com
Phone: +1-604-555-0188

PICKUP:
Location: TechVista Distribution Center
Address: 4521 Boundary Road, Burnaby, BC V5R 2N8
Date: July 25, 2026
Time: 08:00 - 12:00

DELIVERY:
Location: Best Buy DC - Ontario
Address: 2155 Dunwin Drive, Mississauga, ON L5L 1X2
Date: July 28, 2026
Time: 06:00 - 18:00

Commodity: Consumer Electronics - TV Displays
Equipment: Dry Van
Weight: 38,500 lbs
Pallets: 20
Freight Type: FTL
Hazmat: No

Special: Fragile cargo - no double stacking.
""")

# --- Scenario 2: Missing Fields (Word) ---
print("\n2. Missing Fields (Word)")
create_word("scenario2-steelmax-missing.docx", """
Customer: SteelMax Industries
Contact: Robert Kraft
Email: bharathm@ideyalabs.com
Phone: 416-555-0234

Pickup from: SteelMax Fabrication Yard
Address: 89 Industrial Parkway, Hamilton, ON L8W 3B2
Pickup Date: July 28, 2026

Commodity: Structural Steel I-Beams
Weight: 44,000 lbs
Equipment: Flatbed
26 pieces, strapped and blocked

NOTE: Delivery details TBD - customer will confirm.
""")

# --- Scenario 3: Hazmat (PDF) ---
print("\n3. Hazmat Order (PDF)")
create_pdf("scenario3-chemflow-hazmat.pdf", """
HAZARDOUS MATERIALS SHIPMENT REQUEST

Customer: ChemFlow Corporation
Contact: Dr. Sarah Williams
Email: bharathm@ideyalabs.com
Phone: +1-905-555-0199

PICKUP:
Location: ChemFlow Manufacturing Plant
Address: 1200 Chemical Lane, Sarnia, ON N7T 7H5
Date: July 30, 2026
Time: 06:00 - 10:00

DELIVERY:
Location: Dow Chemical Receiving
Address: 5555 Petrochemical Drive, Fort Saskatchewan, AB T8L 4G2
Date: August 2, 2026
Time: 08:00 - 16:00

Commodity: Industrial Solvents - Toluene
Equipment: Tanker
Weight: 42,000 lbs
Freight Type: FTL

*** HAZMAT ***
UN Number: UN1294
Hazmat Class: 3 (Flammable Liquid)
Packing Group: II

Driver Requirements: HAZMAT endorsement, WHMIS training
Emergency Contact: 1-800-CHEMFLOW
""")

# --- Scenario 4: Reefer (Image/PNG) ---
print("\n4. Reefer Order (PNG)")
create_image("scenario4-seafood-reefer.png", """
REFRIGERATED TRANSPORT REQUEST

Customer: Northern Seafood Distributors
Contact: James MacLeod
Email: bharathm@ideyalabs.com
Phone: 902-555-0167

FROM:
Location: Northern Seafood Processing
Address: 45 Harbour Drive, Lunenburg, NS B0J 2C0
Date: July 22, 2026
Time: 05:00 - 07:00

TO:
Location: Metro Inc. Distribution Center
Address: 11011 Blvd Maurice-Duplessis, Montreal, QC H1C 1V6
Date: July 24, 2026
Time: 04:00 - 10:00

Commodity: Fresh Atlantic Lobster & Shrimp
Equipment: Reefer
Temperature: 2C to 4C (MUST NOT EXCEED 4C)
Weight: 28,000 lbs
Pallets: 14
Freight: FTL
Hazmat: No
""")

# --- Scenario 5: LTL Multi-Order (Excel) ---
print("\n5. LTL Orders (Excel - multiple rows)")
create_excel("scenario5-multi-order.xlsx", [
    {
        "Customer": "Cascade Paper Products",
        "Contact": "Lisa Wong",
        "Email": "bharathm@ideyalabs.com",
        "Phone": "604-555-0145",
        "Pickup Location": "Cascade Mill #3",
        "Pickup Address": "777 River Road, Prince George, BC V2L 5R5",
        "Pickup Date": "2026-07-26",
        "Delivery Location": "Staples Canada Warehouse",
        "Delivery Address": "500 Consumers Road, North York, ON M2J 1P8",
        "Delivery Date": "2026-07-30",
        "Commodity": "Copy Paper & Cardstock",
        "Equipment": "Dry Van",
        "Weight (lbs)": 8500,
        "Pallets": 6,
        "Freight Type": "LTL",
        "Hazmat": "No",
    },
    {
        "Customer": "Alpine Dairy Co-op",
        "Contact": "Pierre Dubois",
        "Email": "bharathm@ideyalabs.com",
        "Phone": "418-555-0123",
        "Pickup Location": "Alpine Dairy Processing",
        "Pickup Address": "234 Chemin du Lait, Victoriaville, QC G6P 4S3",
        "Pickup Date": "2026-07-27",
        "Delivery Location": "Sobeys Distribution",
        "Delivery Address": "1500 Lake City Way, Burnaby, BC V5A 4N6",
        "Delivery Date": "2026-08-01",
        "Commodity": "Cheese Products (Refrigerated)",
        "Equipment": "Reefer",
        "Weight (lbs)": 32000,
        "Pallets": 18,
        "Freight Type": "FTL",
        "Hazmat": "No",
    },
])

# --- Scenario 6: Rate Confirmation (PDF) ---
print("\n6. Rate Confirmation (PDF)")
create_pdf("scenario6-rate-confirmation.pdf", """
RATE CONFIRMATION / LOAD TENDER

Broker: Pacific Coast Logistics
Broker Contact: Amanda Torres
Email: bharathm@ideyalabs.com
Phone: 778-555-0156
Reference #: PCL-2026-4521

SHIPPER:
Company: Canfor Corporation
Location: Canfor Pulp Mill
Address: 1500 Pulp Mill Road, Prince George, BC V2N 4W7
Date: July 29, 2026
Time: 07:00 - 11:00
Contact: Mill Dispatch 250-555-0100

CONSIGNEE:
Company: Georgia-Pacific Receiving
Location: GP Distribution Center
Address: 3400 Industrial Blvd, Hinton, AB T7V 1V3
Date: July 30, 2026
Time: 08:00 - 16:00
Contact: Receiving Dock 780-555-0200

LOAD DETAILS:
Commodity: Northern Bleached Softwood Kraft Pulp
Weight: 45,000 lbs
Pieces: 22 unitized bales
Equipment: Flatbed (tarped)
Freight Type: FTL
Hazmat: No

RATE: $3,250.00 all-in (includes FSC)
Payment Terms: Net 30

Driver Requirements:
- Flatbed experience required
- Tarps and straps provided by shipper
- No overweight - scale at mill exit
""")

print("\n" + "=" * 50)
print("All fixtures generated in: test-emails/fixtures/")
print("Files ready to attach to emails for testing.")
